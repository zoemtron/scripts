import pandas as pd
import openpyxl
from openpyxl.styles import Border, Side, Alignment, PatternFill, Font

# 1. Зареждане на файловете
bank = pd.read_excel('bank.xls', header=9)
inv1 = pd.read_excel('invoices1.xls', header=2)
inv2 = pd.read_excel('invoices2.xls', header=2)

# Обединяване на фактурите
invoices = pd.concat([inv1, inv2], ignore_index=True)

# Почистване на имената на колоните
bank.columns = bank.columns.astype(str).str.strip()
invoices.columns = invoices.columns.astype(str).str.strip()

# Дефиниране на имената на колоните
bank_partner_col = "Получател/Наредител:"
bank_credit_col = "Кредит:"
bank_date_col = "Счет. дата:"
bank_ref_col = "Основание за плащане:"

inv_num_col = "№"
inv_partner_col = "Партньор"
inv_amount_col = "Сума за плащане"
inv_date_col = "Дата"

details_col_name = "Банково плащане (Фирма | Дата | Сума | Основание)"

def format_date(val):
    if pd.isna(val) or not val:
        return ''
    try:
        dt = pd.to_datetime(val, dayfirst=True, errors='coerce')
        if pd.notna(dt):
            return dt.strftime('%d-%m-%Y')
        return str(val).split(' ')[0]
    except:
        return str(val).split(' ')[0]

# Извличане на дата за името на изходния файл
valid_dates = pd.to_datetime(invoices[inv_date_col], dayfirst=True, errors='coerce').dropna()
date_suffix = valid_dates.iloc[0].strftime('%m_%Y') if not valid_dates.empty else '01_2026'
output_filename = f'Обработени_Фактури_{date_suffix}.xlsx'

# Филтриране на банковите плащания
bank_credits = bank.dropna(subset=[bank_credit_col]).copy()

# 2. Филтриране и изчистване на фактурите
valid_invoices = []
for _, inv in invoices.iterrows():
    partner_raw = inv.get(inv_partner_col, '')
    partner = str(partner_raw).strip() if pd.notna(partner_raw) else ''
    amount = inv.get(inv_amount_col)
    
    if not partner or partner.lower() == 'общо' or pd.isna(amount):
        continue
        
    inv_num = inv.get(inv_num_col, '')
    inv_num = '' if pd.isna(inv_num) else str(inv_num).split('.')[0].strip()
    inv_date = format_date(inv.get(inv_date_col, ''))
    
    valid_invoices.append({
        'Фактура': inv_num,
        'Дата': inv_date,
        'Фирма': partner,
        'Сума за плащане': amount
    })

inv_df = pd.DataFrame(valid_invoices)

processed_rows = []
used_bank_indices = set()
unmatched_invoices = []

# 3. ЕТАП 1: Търсене на ТОЧНО съвпадение (Фирма + Сума)
for idx, row in inv_df.iterrows():
    partner = row['Фирма']
    amount = row['Сума за плащане']
    
    available_bank = bank_credits[~bank_credits.index.isin(used_bank_indices)]
    bank_partners_clean = available_bank[bank_partner_col].astype(str).str.strip()
    
    exact_match = available_bank[
        (bank_partners_clean == partner) & 
        (available_bank[bank_credit_col] == amount)
    ]
    
    if not exact_match.empty:
        matched_idx = exact_match.index[0]
        used_bank_indices.add(matched_idx)
        
        b_row = exact_match.iloc[0]
        b_partner = str(b_row[bank_partner_col]).strip() if pd.notna(b_row[bank_partner_col]) else partner
        b_date = format_date(b_row[bank_date_col])
        b_amount = b_row[bank_credit_col]
        b_ref = str(b_row[bank_ref_col]).strip() if pd.notna(b_row[bank_ref_col]) else ''
        
        details = f"{b_partner}\nДата: {b_date} | Сума: {b_amount} | Основание: {b_ref}"
        
        processed_rows.append({
            'Фактура': row['Фактура'],
            'Дата': row['Дата'],
            'Фирма': partner,
            'Сума за плащане': amount,
            'Статус': 'ПЛАТЕНА',
            details_col_name: details
        })
    else:
        unmatched_invoices.append(row)

# 4. ЕТАП 2: Търсене на "ПРОВЕРИ" за останалите
remaining_bank = bank_credits[~bank_credits.index.isin(used_bank_indices)]

for row in unmatched_invoices:
    inv_num = row['Фактура']
    partner = row['Фирма']
    amount = row['Сума за плащане']
    
    ref_match = remaining_bank[
        remaining_bank[bank_ref_col].astype(str).str.contains(inv_num, regex=False)
    ] if inv_num else pd.DataFrame()
    
    matched_b_row = None
    
    if len(ref_match) == 1:
        matched_b_row = ref_match.iloc[0]
    else:
        matching_amount_rows = remaining_bank[remaining_bank[bank_credit_col] == amount]
        competing_inv_count = sum(1 for item in unmatched_invoices if item['Сума за плащане'] == amount)
        
        if len(matching_amount_rows) == 1 and competing_inv_count == 1:
            matched_b_row = matching_amount_rows.iloc[0]
            
    if matched_b_row is not None:
        b_partner = str(matched_b_row[bank_partner_col]).strip() if pd.notna(matched_b_row[bank_partner_col]) else ''
        b_date = format_date(matched_b_row[bank_date_col])
        b_amount = matched_b_row[bank_credit_col]
        b_ref = str(matched_b_row[bank_ref_col]).strip() if pd.notna(matched_b_row[bank_ref_col]) else ''
        
        details = f"{b_partner}\nДата: {b_date} | Сума: {b_amount} | Основание: {b_ref}"
        status = 'ПРОВЕРИ'
    else:
        details = ''
        status = 'НЕПЛАТЕНА'
        
    processed_rows.append({
        'Фактура': inv_num,
        'Дата': row['Дата'],
        'Фирма': partner,
        'Сума за плащане': amount,
        'Статус': status,
        details_col_name: details
    })

# Формиране на DataFrames
cols_order = ['Фактура', 'Дата', 'Фирма', 'Сума за плащане', 'Статус', details_col_name]
all_df = pd.DataFrame(processed_rows)[cols_order]

# Филтриране и сортиране по ДАТА за Sheet 1 ("Неплатени")
unpaid_df = all_df[all_df['Статус'].isin(['ПРОВЕРИ', 'НЕПЛАТЕНА'])].copy()
unpaid_df['_sort_date'] = pd.to_datetime(unpaid_df['Дата'], format='%d-%m-%Y', errors='coerce')
unpaid_df = unpaid_df.sort_values(by='_sort_date').drop(columns=['_sort_date'])

# 5. Записване с 2 Sheet-а и форматиране
def save_excel_with_two_sheets(df_unpaid, df_all, filename):
    black_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    fills = {
        'ПЛАТЕНА': PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid'),
        'ПРОВЕРИ': PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid'),
        'НЕПЛАТЕНА': PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid')
    }
    
    fonts = {
        'ПЛАТЕНА': Font(color='276A3C', bold=True),
        'ПРОВЕРИ': Font(color='B25900', bold=True),
        'НЕПЛАТЕНА': Font(color='C00000', bold=True)
    }

    max_widths = {
        'Фирма': 30,
        details_col_name: 45
    }

    sheets_data = [
        ('Неплатени', df_unpaid),
        ('Всички', df_all)
    ]

    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        for sheet_name, df in sheets_data:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]
            
            if df.empty:
                continue

            status_col_idx = df.columns.get_loc('Статус') + 1

            for row_idx in range(2, len(df) + 2):
                status_cell = worksheet.cell(row=row_idx, column=status_col_idx)
                val = str(status_cell.value).strip()
                if val in fills:
                    status_cell.fill = fills[val]
                    status_cell.font = fonts[val]
                    status_cell.alignment = Alignment(horizontal='center', vertical='center')

            for col_idx, col in enumerate(worksheet.columns, start=1):
                col_name = df.columns[col_idx - 1]
                max_len = 0
                
                for cell in col:
                    cell.border = black_border
                    val_str = str(cell.value or '')
                    
                    if col_name in max_widths or '\n' in val_str:
                        cell.alignment = Alignment(wrap_text=True, vertical='center')
                        
                    for line in val_str.split('\n'):
                        if len(line) > max_len:
                            max_len = len(line)
                
                calculated_width = max(max_len + 3, 12)
                if col_name in max_widths:
                    final_width = min(calculated_width, max_widths[col_name])
                else:
                    final_width = calculated_width
                    
                col_letter = openpyxl.utils.get_column_letter(col_idx)
                worksheet.column_dimensions[col_letter].width = final_width

save_excel_with_two_sheets(unpaid_df, all_df, output_filename)

print(f"Готово! Файлът '{output_filename}' е генериран със сортиран по дата шийт 'Неплатени'.")